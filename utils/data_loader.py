"""
测试数据加载工具 (数据驱动)
- 支持 yaml / json / excel
- 用例通过 load_yaml("login_data.yaml") 拿到参数列表，配合 @pytest.mark.parametrize
"""

import json

import yaml
from openpyxl import load_workbook

from config.settings import settings

DATA_DIR = settings.root_dir / "data"


def load_yaml(filename: str) -> list | dict:
    path = DATA_DIR / filename
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def load_json(filename: str) -> list | dict:
    path = DATA_DIR / filename
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def load_excel(filename: str, sheet: str = None) -> list[dict]:
    """读取 excel，首行作为表头，返回 [{列名: 值}, ...]"""
    path = DATA_DIR / filename
    wb = load_workbook(path, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows or not any(header is not None for header in rows[0]):
            raise ValueError(f"Excel 工作表没有表头: {path} | sheet={ws.title}")
        headers = rows[0]
        return [dict(zip(headers, row, strict=False)) for row in rows[1:]]
    finally:
        wb.close()


def read_lines(filename: str) -> list[str]:
    """
    读取纯文本文件，每行一个值（如门店编码列表）。
    - 自动去除空白行和以 # 开头的注释行
    - filename 相对 data/ 目录，如 "sales/m餐厅.txt"
    - 文件不存在时返回 []（不报错，便于 CI 上缺数据时跳过）
    """
    path = DATA_DIR / filename
    if not path.exists():
        from utils.logger import log

        log.warning(f"数据文件不存在，返回空列表: {path}")
        return []
    result = []
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith("#"):
                result.append(line)
    return result
